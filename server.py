"""Two-stage FastAPI application for multi-PO purchase-order plan sheets."""

from __future__ import annotations

import json
import re
import shutil
import tempfile
import uuid
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook

from artwork_archive import ArchiveError
from artwork_package import resolve_artwork_package
from extract_to_plan_sheet import (
    FIRST_DATA_ROW,
    WORKBOOK_FORMATS,
    add_artwork_to_row,
    build_workbook,
    extract_purchase_order,
)

APP_DIR = Path(__file__).resolve().parent
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_PO_FILES = 25
ZIP_MAGIC_BYTES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
sessions: dict[str, Path] = {}

app = FastAPI(title="PDF to Plan Sheet", docs_url=None, redoc_url=None)
app.mount("/static", StaticFiles(directory=APP_DIR / "static"), name="static")


def safe_name(value: str, suffix: str) -> str:
    name = re.sub(r"[^A-Za-z0-9 ._+-]", "_", Path(value.strip()).name).strip(" .") or "plan_sheet"
    return (name if name.lower().endswith(suffix) else name + suffix)[:180]


async def save_upload(upload: UploadFile, directory: Path, name: str, magic_bytes: tuple[bytes, ...], kind: str) -> Path:
    if not upload.filename:
        raise HTTPException(status_code=400, detail=f"Every {kind} upload needs a filename.")
    contents = await upload.read()
    if not any(contents.startswith(magic) for magic in magic_bytes):
        raise HTTPException(status_code=400, detail=f"{upload.filename} is not a valid {kind} file.")
    if len(contents) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"{upload.filename} exceeds the 25 MB limit.")
    path = directory / name
    path.write_bytes(contents)
    return path


async def save_pdf(upload: UploadFile, directory: Path, name: str) -> Path:
    return await save_upload(upload, directory, name, (b"%PDF",), "PDF")


def get_session(session_id: str) -> Path:
    directory = sessions.get(session_id)
    if not directory or not directory.is_dir():
        raise HTTPException(status_code=404, detail="This workbook session has expired. Create the plan sheet again.")
    return directory


@app.get("/", include_in_schema=False)
def homepage() -> FileResponse:
    return FileResponse(APP_DIR / "static" / "index.html", media_type="text/html")


@app.get("/health", include_in_schema=False)
def health() -> dict[str, str]:
    """Liveness/readiness check for the host (e.g. Render) -- cheap, no session state involved."""
    missing = [fmt.label for fmt in WORKBOOK_FORMATS.values() if not fmt.template_path.is_file()]
    if missing:
        raise HTTPException(status_code=503, detail=f"Missing workbook blueprint(s): {', '.join(missing)}.")
    return {"status": "ok"}


def default_workbook_name(sheet_titles: list[str]) -> str:
    label = " + ".join(sheet_titles) if len(sheet_titles) <= 3 else f"{sheet_titles[0]} + {len(sheet_titles) - 1} more"
    return safe_name(f"Plan Sheet - {label}", ".xlsx")


@app.post("/api/convert")
async def convert_pdf(format: str = Form(...), pdfs: list[UploadFile] = File(...)) -> dict[str, object]:
    fmt = WORKBOOK_FORMATS.get(format)
    if fmt is None:
        raise HTTPException(status_code=400, detail=f"Unknown workbook format {format!r}.")
    if not fmt.template_path.is_file():
        raise HTTPException(status_code=500, detail=f"The {fmt.label} workbook blueprint is missing from the server.")
    if not pdfs:
        raise HTTPException(status_code=400, detail="Add at least one purchase-order PDF.")
    if len(pdfs) > MAX_PO_FILES:
        raise HTTPException(status_code=413, detail=f"Add at most {MAX_PO_FILES} purchase-order PDFs at once.")
    session_id = uuid.uuid4().hex
    workdir = Path(tempfile.mkdtemp(prefix="plan-sheet-"))
    try:
        purchase_orders: list[tuple[str, list[dict]]] = []
        seen_po_numbers: dict[str, str] = {}
        for index, upload in enumerate(pdfs):
            source = await save_pdf(upload, workdir, f"po_{index}.pdf")
            po_number, rows = extract_purchase_order(source)
            if po_number in seen_po_numbers:
                raise HTTPException(
                    status_code=422,
                    detail=(
                        f"Purchase Order {po_number} showed up twice — in "
                        f"{seen_po_numbers[po_number]} and {upload.filename}. "
                        "Combine them into one PDF, or remove one, then try again."
                    ),
                )
            seen_po_numbers[po_number] = upload.filename or source.name
            purchase_orders.append((po_number, rows))
        output = workdir / "plan_sheet.xlsx"
        result = build_workbook(fmt, output, purchase_orders)
        sheet_rows = result["sheets"]
        workbook_name = default_workbook_name(list(sheet_rows.keys()))
        (workdir / "meta.json").write_text(
            json.dumps({
                "name": workbook_name,
                "sheets": sheet_rows,
                "format": fmt.key,
            }),
            encoding="utf-8",
        )
        sessions[session_id] = workdir
    except HTTPException:
        shutil.rmtree(workdir, ignore_errors=True)
        raise
    except RuntimeError as exc:
        shutil.rmtree(workdir, ignore_errors=True)
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except Exception as exc:
        shutil.rmtree(workdir, ignore_errors=True)
        raise HTTPException(status_code=500, detail="Could not create the Excel workbook.") from exc
    return {
        "session_id": session_id,
        "download_url": f"/api/session/{session_id}/workbook",
        "workbook_name": workbook_name,
        "sheets": [
            {"title": title, "po_number": po_number, "row_count": len(rows)}
            for (po_number, rows), title in zip(purchase_orders, sheet_rows.keys())
        ],
        "row_count": sum(len(rows) for rows in sheet_rows.values()),
    }


@app.post("/api/enrich")
async def enrich_workbook(session_id: str = Form(...), package: UploadFile = File(...)) -> dict[str, object]:
    workdir = get_session(session_id)
    meta = json.loads((workdir / "meta.json").read_text(encoding="utf-8"))
    fmt = WORKBOOK_FORMATS[meta["format"]]
    known_contracts = {
        str(row["contract"]).upper()
        for rows in meta["sheets"].values()
        for row in rows
        if row.get("contract")
    }
    contents = await package.read()
    filename = package.filename or ""
    if not filename.lower().endswith(".zip") or not contents.startswith(ZIP_MAGIC_BYTES):
        raise HTTPException(status_code=400, detail="Upload one valid ZIP artwork package.")
    try:
        resolution = resolve_artwork_package(contents, fmt.key, known_contracts)
    except ArchiveError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    workbook = load_workbook(workdir / "plan_sheet.xlsx")
    for sheet in workbook.worksheets:
        for row in range(FIRST_DATA_ROW, sheet.max_row):
            contract = sheet.cell(row, 2).value
            if not contract:
                continue
            contract = str(contract).upper()
            for slot in fmt.image_slots:
                image_bytes = resolution.slots.get((contract, slot.source))
                if image_bytes:
                    add_artwork_to_row(sheet, row, slot.column, image_bytes, slot.box)
    workbook.save(workdir / "completed_plan_sheet.xlsx")
    records = [record.to_dict() for record in resolution.records]
    return {
        "download_url": f"/api/session/{session_id}/completed",
        "report": {"records": records, "ignored": list(resolution.ignored)},
        "matched_count": sum(record["status"] in {"matched", "fallback"} for record in records),
        "fallback_count": sum(record["status"] == "fallback" for record in records),
        "slot_count": len(records),
    }


@app.get("/api/session/{session_id}/workbook")
def download_workbook(session_id: str, name: Optional[str] = None) -> FileResponse:
    workdir = get_session(session_id)
    meta = json.loads((workdir / "meta.json").read_text(encoding="utf-8"))
    filename = safe_name(name, ".xlsx") if name else meta["name"]
    return FileResponse(workdir / "plan_sheet.xlsx", filename=filename, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/session/{session_id}/completed")
def download_completed(session_id: str, name: Optional[str] = None) -> FileResponse:
    workdir = get_session(session_id)
    path = workdir / "completed_plan_sheet.xlsx"
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Add artwork before downloading the completed workbook.")
    meta = json.loads((workdir / "meta.json").read_text(encoding="utf-8"))
    filename = safe_name(name, ".xlsx") if name else meta["name"]
    return FileResponse(path, filename=filename, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
